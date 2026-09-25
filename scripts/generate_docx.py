#!/usr/bin/env python3
"""
generate_docx.py - Generates an editable Word (.docx) resume with Claude frontend styling
Part of Antigravity skill: resume-claude-stylist
"""

import os
import sys
import json
import argparse
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

COLOR_PRIMARY = RGBColor(0x1F, 0x1D, 0x1A)      # Charcoal #1F1D1A
COLOR_TERRACOTTA = RGBColor(0xB2, 0x4C, 0x2C)   # Claude Terracotta #B24C2C
COLOR_MUTED = RGBColor(0x7A, 0x74, 0x6B)        # Muted gray #7A746B
COLOR_BODY = RGBColor(0x2D, 0x2A, 0x26)         # Body text #2D2A26
COLOR_TAG = RGBColor(0x55, 0x50, 0x4A)          # Tag text

def set_run_font(run, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.8, bold=False, italic=False, color=COLOR_BODY):
    run.font.name = font_en
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_en)
    rFonts.set(qn('w:hAnsi'), font_en)
    rFonts.set(qn('w:eastAsia'), font_cn)
    rFonts.set(qn('w:cs'), font_en)
    rPr.append(rFonts)

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_section_header(doc, title_cn, title_en):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_bar = p.add_run('▍ ')
    set_run_font(r_bar, font_en='Segoe UI', font_cn='宋体', size_pt=10.5, bold=True, color=COLOR_TERRACOTTA)
    
    r_title = p.add_run(title_cn)
    set_run_font(r_title, font_en='Georgia', font_cn='宋体', size_pt=10.5, bold=True, color=COLOR_PRIMARY)
    
    r_sub = p.add_run(f'  {title_en}')
    set_run_font(r_sub, font_en='Segoe UI', font_cn='微软雅黑', size_pt=6.8, bold=True, color=COLOR_MUTED)
    
    pPr = p._element.get_or_add_pPr()
    pBdr = parse_xml(r'''
        <w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:bottom w:val="single" w:sz="4" w:space="2" w:color="E6E1D6"/>
        </w:pBdr>
    ''')
    pPr.append(pBdr)

def build_docx_from_data(data, output_path, avatar_path=None):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Cm(1.15)
        section.bottom_margin = Cm(1.0)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # 1. Header Table
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    left_cell = header_table.cell(0, 0)
    left_cell.width = Cm(15.2)
    right_cell = header_table.cell(0, 1)
    right_cell.width = Cm(2.8)

    set_cell_margins(left_cell, top=0, bottom=30, left=0, right=100)
    set_cell_margins(right_cell, top=0, bottom=30, left=50, right=0)

    # Name & Target
    p_name = left_cell.paragraphs[0]
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.line_spacing = 1.1

    r_name = p_name.add_run(f"{data.get('name', '候选人')}   ")
    set_run_font(r_name, font_en='Georgia', font_cn='宋体', size_pt=20, bold=True, color=COLOR_PRIMARY)

    if data.get('target_role'):
        r_badge = p_name.add_run(f"求职意向：{data['target_role']}")
        set_run_font(r_badge, font_en='Segoe UI', font_cn='微软雅黑', size_pt=9.5, bold=True, color=COLOR_TERRACOTTA)

    # Demographics
    demographics = data.get('demographics', [])
    if demographics:
        p_meta1 = left_cell.add_paragraph()
        p_meta1.paragraph_format.space_before = Pt(1)
        p_meta1.paragraph_format.space_after = Pt(1)
        p_meta1.paragraph_format.line_spacing = 1.15
        for idx, item in enumerate(demographics):
            r_l = p_meta1.add_run(item['label'] + '：')
            set_run_font(r_l, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, bold=True, color=COLOR_PRIMARY)
            r_v = p_meta1.add_run(item['value'])
            set_run_font(r_v, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, bold=False, color=COLOR_TAG)
            if idx < len(demographics) - 1:
                r_sep = p_meta1.add_run('   /   ')
                set_run_font(r_sep, font_en='Segoe UI', font_cn='微软雅黑', size_pt=7.5, color=COLOR_MUTED)

    # Contacts
    contacts = data.get('contacts', [])
    if contacts:
        p_meta2 = left_cell.add_paragraph()
        p_meta2.paragraph_format.space_before = Pt(1)
        p_meta2.paragraph_format.space_after = Pt(2)
        p_meta2.paragraph_format.line_spacing = 1.15
        for idx, item in enumerate(contacts):
            r_l = p_meta2.add_run(item['label'] + '：')
            set_run_font(r_l, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, bold=True, color=COLOR_PRIMARY)
            r_v = p_meta2.add_run(item['value'])
            set_run_font(r_v, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, bold=False, color=COLOR_TAG)
            if idx < len(contacts) - 1:
                r_sep = p_meta2.add_run('   /   ')
                set_run_font(r_sep, font_en='Segoe UI', font_cn='微软雅黑', size_pt=7.5, color=COLOR_MUTED)

    # Photo
    if avatar_path and os.path.exists(avatar_path):
        p_photo = right_cell.paragraphs[0]
        p_photo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_photo.paragraph_format.space_before = Pt(0)
        p_photo.paragraph_format.space_after = Pt(0)
        r_img = p_photo.add_run()
        r_img.add_picture(avatar_path, width=Cm(2.0), height=Cm(2.67))

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(2)
    pPr = p_div._element.get_or_add_pPr()
    pBdr = parse_xml(r'''
        <w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:bottom w:val="single" w:sz="6" w:space="1" w:color="EAE5DC"/>
        </w:pBdr>
    ''')
    pPr.append(pBdr)

    # 2. Education
    if data.get('education'):
        add_section_header(doc, '教育背景', 'EDUCATION')
        for edu in data['education']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            r_s = p.add_run(edu['school'])
            set_run_font(r_s, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.8, bold=True, color=COLOR_PRIMARY)
            r_sep = p.add_run(' ｜ ')
            set_run_font(r_sep, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, color=COLOR_MUTED)
            r_d = p.add_run(edu['major'])
            set_run_font(r_d, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.8, bold=False, color=COLOR_BODY)
            r_tab = p.add_run(f"\t{edu['period']}")
            set_run_font(r_tab, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, color=COLOR_MUTED)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)

    # 3. Evaluation
    if data.get('evaluation'):
        add_section_header(doc, '自我评价', 'ABOUT ME')
        for item in data['evaluation']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.35)
            r_bullet = p.add_run('• ')
            set_run_font(r_bullet, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.4, color=COLOR_TERRACOTTA)
            r_t = p.add_run(item['title'] + '：')
            set_run_font(r_t, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.4, bold=True, color=COLOR_PRIMARY)
            r_b = p.add_run(item['detail'])
            set_run_font(r_b, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.4, color=COLOR_BODY)

    # 4. Skills
    if data.get('skills'):
        add_section_header(doc, '专业技能', 'SKILLS')
        for item in data['skills']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.35)
            r_l = p.add_run(item['category'] + '：')
            set_run_font(r_l, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.4, bold=True, color=COLOR_TERRACOTTA)
            r_d = p.add_run(item['description'])
            set_run_font(r_d, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.4, color=COLOR_BODY)

    # 5. Internship
    if data.get('internships'):
        add_section_header(doc, '实习经历', 'EXPERIENCE')
        for exp in data['internships']:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(1)
            p_exp.paragraph_format.space_after = Pt(1)
            p_exp.paragraph_format.line_spacing = 1.15
            r_c = p_exp.add_run(exp['company'])
            set_run_font(r_c, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.8, bold=True, color=COLOR_PRIMARY)
            r_sep = p_exp.add_run(' ｜ ')
            set_run_font(r_sep, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, color=COLOR_MUTED)
            r_r = p_exp.add_run(exp['role'])
            set_run_font(r_r, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.5, bold=True, color=COLOR_BODY)
            r_tab = p_exp.add_run(f"\t{exp['period']}")
            set_run_font(r_tab, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, color=COLOR_MUTED)
            p_exp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
            for b in exp.get('bullets', []):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.left_indent = Cm(0.35)
                r_bullet = p.add_run('▪ ')
                set_run_font(r_bullet, font_en='Segoe UI', font_cn='微软雅黑', size_pt=7, color=COLOR_TERRACOTTA)
                r_b = p.add_run(b)
                set_run_font(r_b, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, color=COLOR_BODY)

    # 6. Projects
    if data.get('projects'):
        add_section_header(doc, '项目经历', 'PROJECTS')
        for proj in data['projects']:
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(1.5)
            p_h.paragraph_format.space_after = Pt(0.5)
            p_h.paragraph_format.line_spacing = 1.15
            r_t = p_h.add_run(proj['title'])
            set_run_font(r_t, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.8, bold=True, color=COLOR_PRIMARY)
            r_tab = p_h.add_run(f"\t{proj['period']}")
            set_run_font(r_tab, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, color=COLOR_MUTED)
            p_h.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
            
            if proj.get('tech_stack'):
                p_s = doc.add_paragraph()
                p_s.paragraph_format.space_before = Pt(0)
                p_s.paragraph_format.space_after = Pt(1)
                p_s.paragraph_format.line_spacing = 1.15
                p_s.paragraph_format.left_indent = Cm(0.35)
                r_sl = p_s.add_run('技术栈：')
                set_run_font(r_sl, font_en='Segoe UI', font_cn='微软雅黑', size_pt=7.8, bold=True, color=COLOR_MUTED)
                r_sv = p_s.add_run(proj['tech_stack'])
                set_run_font(r_sv, font_en='Segoe UI', font_cn='微软雅黑', size_pt=7.8, color=COLOR_TAG)
                
            for b in proj.get('bullets', []):
                p_b = doc.add_paragraph()
                p_b.paragraph_format.space_before = Pt(0)
                p_b.paragraph_format.space_after = Pt(1)
                p_b.paragraph_format.line_spacing = 1.15
                p_b.paragraph_format.left_indent = Cm(0.35)
                r_bullet = p_b.add_run('▪ ')
                set_run_font(r_bullet, font_en='Segoe UI', font_cn='微软雅黑', size_pt=6.8, color=COLOR_TERRACOTTA)
                parts = b.split('【')
                for i, pt in enumerate(parts):
                    if i == 0:
                        r_part = p_b.add_run(pt)
                        set_run_font(r_part, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, color=COLOR_BODY)
                    else:
                        subparts = pt.split('】')
                        r_bold = p_b.add_run(subparts[0])
                        set_run_font(r_bold, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, bold=True, color=COLOR_PRIMARY)
                        if len(subparts) > 1:
                            r_rest = p_b.add_run(subparts[1])
                            set_run_font(r_rest, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, color=COLOR_BODY)
                            
            if proj.get('result'):
                p_res = doc.add_paragraph()
                p_res.paragraph_format.space_before = Pt(0)
                p_res.paragraph_format.space_after = Pt(1)
                p_res.paragraph_format.line_spacing = 1.15
                p_res.paragraph_format.left_indent = Cm(0.35)
                r_bullet = p_res.add_run('• ')
                set_run_font(r_bullet, font_en='Segoe UI', font_cn='微软雅黑', size_pt=7.5, bold=True, color=COLOR_TERRACOTTA)
                r_rlbl = p_res.add_run('项目成效：')
                set_run_font(r_rlbl, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, bold=True, color=COLOR_TERRACOTTA)
                parts = proj['result'].split('【')
                for i, pt in enumerate(parts):
                    if i == 0:
                        r_part = p_res.add_run(pt)
                        set_run_font(r_part, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, color=COLOR_BODY)
                    else:
                        subparts = pt.split('】')
                        r_bold = p_res.add_run(subparts[0])
                        set_run_font(r_bold, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, bold=True, color=COLOR_PRIMARY)
                        if len(subparts) > 1:
                            r_rest = p_res.add_run(subparts[1])
                            set_run_font(r_rest, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, color=COLOR_BODY)

    # 7. Certifications
    if data.get('certifications'):
        add_section_header(doc, '技能证书', 'CERTIFICATIONS')
        p_cert = doc.add_paragraph()
        p_cert.paragraph_format.space_before = Pt(1)
        p_cert.paragraph_format.space_after = Pt(0)
        p_cert.paragraph_format.line_spacing = 1.15
        p_cert.paragraph_format.left_indent = Cm(0.35)
        for cert in data['certifications']:
            r_c = p_cert.add_run(f"［ {cert} ］   ")
            set_run_font(r_c, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.3, bold=True, color=COLOR_PRIMARY)
        if data.get('extra_cert_note'):
            r_note = p_cert.add_run(data['extra_cert_note'])
            set_run_font(r_note, font_en='Segoe UI', font_cn='微软雅黑', size_pt=8.0, color=COLOR_MUTED)

    doc.save(output_path)
    print(f"[SUCCESS] Exported Word Resume: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Generate Word (.docx) resume from JSON data.")
    parser.add_argument("--data", required=True, help="Path to JSON resume data")
    parser.add_argument("--output", required=True, help="Output .docx path")
    parser.add_argument("--avatar", help="Optional avatar image path")
    args = parser.parse_args()

    with open(args.data, 'r', encoding='utf-8') as f:
        data = json.load(f)

    build_docx_from_data(data, args.output, args.avatar)

if __name__ == "__main__":
    main()
